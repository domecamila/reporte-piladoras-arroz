"""Reporte de piladoras desde ArcGIS Survey123. Sin HTML ni credenciales si la capa es pública."""
import re
import unicodedata
from datetime import datetime
from pathlib import Path
import requests
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

URL = ('https://services.arcgis.com/53Yac7UQ7RmeEyAd/arcgis/rest/services/'
       'service_d007ae8e06fd47e5ab2f49abb6d8d91b/FeatureServer/0')
OUT = Path(__file__).resolve().parent / 'Reportes'
OUT.mkdir(exist_ok=True)
DESTINO = OUT / 'Reporte_Piladoras_Arroz.xlsx'
SESSION = requests.Session()


def get(url, **params):
    r = SESSION.get(url, params={'f': 'json', **params}, timeout=90)
    r.raise_for_status()
    d = r.json()
    if 'error' in d:
        raise RuntimeError(f"Error ArcGIS: {d['error']}")
    return d


def normalizar(x):
    s = unicodedata.normalize('NFD', str(x or '').strip().lower())
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn')


def obtener_campo(campos, nombres=(), prefijo=None, obligatorio=True):
    for n in nombres:
        for c in campos:
            if normalizar(c['name']) == normalizar(n):
                return c['name']
    if prefijo:
        patron = re.compile(r'^' + re.escape(normalizar(prefijo).strip()) + r'(?=\s|$)')
        encontrados = [c['name'] for c in campos if patron.search(normalizar(c.get('alias', '')))]
        if len(encontrados) == 1:
            return encontrados[0]
    if obligatorio:
        raise ValueError(f'No se pudo identificar campo {nombres} / {prefijo}. '\
                         'Revisar los nombres impresos al inicio.')
    return None


def buscar_capacidad(campos):
    # Se buscan tanto los nombres internos como los alias de la pregunta.
    candidatos = []
    for c in campos:
        texto = normalizar(c['name'] + ' ' + c.get('alias', ''))
        if 'capacidad' in texto and ('cascara' in texto or 'arroz' in texto):
            candidatos.append(c['name'])
    if len(candidatos) == 1:
        return candidatos[0]
    if len(candidatos) > 1:
        raise ValueError(f'Capacidad ambigua: {candidatos}. Revisar nombres de la capa.')
    return obtener_campo(
        campos,
        ['capacidad_arroz_cascara', 'capacidad_almacenamiento'],
        obligatorio=False
    )


def valor(fila, campo):
    return fila.get(campo) if campo else None


def coordenadas(registro, C):
    # Primero atributos; si faltan, usar la geometría de la entidad.
    lat = numero(valor(registro, C['latitud']))
    lon = numero(valor(registro, C['longitud']))
    alt = numero(valor(registro, C['altura']))
    geo = registro.get('_geometry') or {}
    gx, gy = numero(geo.get('x')), numero(geo.get('y'))
    if lat is None:
        lat = gy
    if lon is None:
        lon = gx
    if alt is None:
        alt = numero(geo.get('z'))
    # Detectar inversión típica X/Y en Ecuador sin corregir silenciosamente.
    if lat is not None and lon is not None and abs(lat) > 20 and abs(lon) <= 20:
        print(f'AVISO: coordenadas aparentemente invertidas en registro {registro.get("_oid", "")}. Se intercambian.')
        lat, lon = lon, lat
    return lat, lon, alt


def numero(x):
    try:
        if x is None or str(x).strip() == '':
            return None
        return float(str(x).strip().replace(',', '.'))
    except (ValueError, TypeError):
        return None


def peso(x):
    s = normalizar(x)
    s = re.sub(r'^h1_1_1_', '', s)
    m = re.match(r'^(\d+(?:[.,]\d+)?)', s)
    if not m:
        return None
    p = float(m.group(1).replace(',', '.'))
    return p if p in {180, 200, 205, 210, 215, 220, 230} else None


def unidad(x):
    s = normalizar(x)
    if s in {'tonelada', 'toneladas', 't'}:
        return 'toneladas'
    if s in {'saca', 'sacas'}:
        return 'sacas'
    return s


def decodificar(x, campo, metadata):
    if x is None:
        return ''
    s = str(x).strip()
    if not campo:
        return s
    meta = next((c for c in metadata if c['name'] == campo), {})
    domain = (meta.get('domain') or {}).get('codedValues') or []
    mapa = {str(o['code']): str(o['name']) for o in domain}
    if s in mapa:
        return mapa[s]
    tokens = re.split(r'[,;\s]+', s)
    if tokens and all(t in mapa for t in tokens):
        return ', '.join(mapa[t] for t in tokens)
    return s


def tipo_arroz(x):
    s = normalizar(x)
    tokens = set(re.split(r'[,;\s_]+', s))
    cascara = bool(tokens & {'cascara', 'ambas', 'ambos'})
    pilado = bool(tokens & {'pilado', 'ambas', 'ambos'})
    if cascara and pilado:
        return 'Ambas'
    if cascara:
        return 'Cáscara'
    if pilado:
        return 'Pilado'
    return 'No identificado'


def toneladas(cantidad, unidad_original, peso_original, seleccionado):
    if not seleccionado:
        return None, 'No aplica'
    n = numero(cantidad)
    if n is None or n < 0:
        return None, 'Cantidad inválida o vacía'
    u = unidad(unidad_original)
    if u == 'toneladas':
        return n, 'Correcto'
    if u == 'sacas':
        p = peso(peso_original)
        if p is None:
            return None, 'Peso de saca no reconocido'
        return n * p * 0.00045359237, 'Correcto'
    return None, f'Unidad no reconocida: {unidad_original}'


def fecha_ecuador(x):
    if x is None or str(x).strip() == '':
        return pd.NaT
    if isinstance(x, (int, float)):
        return pd.to_datetime(x, unit='ms', utc=True, errors='coerce').tz_convert('America/Guayaquil').tz_localize(None)
    t = pd.to_datetime(x, errors='coerce')
    if pd.isna(t):
        return pd.NaT
    if t.tzinfo is not None:
        return t.tz_convert('America/Guayaquil').tz_localize(None)
    return t


def ejecutar():
    print('Conectando con ArcGIS Survey123...')
    meta = get(URL)
    campos = meta['fields']
    print('Campos disponibles (nombre | alias):')
    for c in campos:
        print(' ', c['name'], '|', c.get('alias', ''))

    # Nombres de campos y alias de la boleta. Ajustables si Survey usa otros nombres.
    C = {
        'provincia': obtener_campo(campos, ['provincia']),
        'canton': obtener_campo(campos, ['canton']),
        'parroquia': obtener_campo(campos, ['parroquia']),
        'latitud': obtener_campo(campos, ['Y', 'latitud', 'latitude'], obligatorio=False),
        'longitud': obtener_campo(campos, ['X', 'longitud', 'longitude'], obligatorio=False),
        'altura': obtener_campo(campos, ['Z', 'altura', 'altitud', 'elevation'], obligatorio=False),
        'ruc': obtener_campo(campos, ['ruc', 'ruc_piladora', 'ruc_organizacion'], prefijo='B1.', obligatorio=False),
        'razon': obtener_campo(campos, ['razon_social', 'razonsocial'], prefijo='B2.', obligatorio=False),
        'piladora': obtener_campo(campos, ['nombre_piladora'], prefijo='B3.'),
        'capacidad': buscar_capacidad(campos),
        'servicios': obtener_campo(campos, ['servicios'], obligatorio=False),
        'tipo': obtener_campo(campos, ['tipo_arroz'], prefijo='C1.'),
        'cant_c': obtener_campo(campos, ['hoy_arroz_cascara'], prefijo='C2.'),
        'unidad_c': obtener_campo(campos, ['medida_hoy_arroz'], prefijo='C2.1'),
        'peso_c': obtener_campo(campos, ['presentacion_saca'], prefijo='C2.1.2', obligatorio=False),
        'cant_p': obtener_campo(campos, ['hoy_arroz_pilado'], prefijo='C3.'),
        'unidad_p': obtener_campo(campos, ['medida_hoy_pilado'], prefijo='C3.1'),
        'peso_p': obtener_campo(campos, ['presentacion_saca_pilado'], prefijo='C3.1.2', obligatorio=False),
        'tecnico': obtener_campo(campos, ['nombre_tecnico', 'tecnico'], obligatorio=False),
        'cedula': obtener_campo(campos, ['cedula_tecnico', 'cedula_del_tecnico'], obligatorio=False),
        'fecha': obtener_campo(campos, ['fecha', 'fecha_levantamiento', 'end'], obligatorio=False),
        'observaciones': obtener_campo(campos, ['observaciones'], obligatorio=False),
    }
    for k, v in C.items():
        print(f'{k:15}: {v}')
    if not C['fecha']:
        raise ValueError('No se identificó el campo fecha; no se puede clasificar levantamientos.')

    oid = meta.get('objectIdField') or next(c['name'] for c in campos if c['type'] == 'esriFieldTypeOID')
    ids = get(URL + '/query', where='1=1', returnIdsOnly='true').get('objectIds') or []
    registros = []
    for i in range(0, len(ids), 200):
        batch = ids[i:i+200]
        d = get(URL + '/query', objectIds=','.join(map(str, batch)), outFields='*', returnGeometry='true', outSR='4326')
        for f in d.get('features', []):
            a = dict(f['attributes'])
            a['_geometry'] = f.get('geometry') or {}
            a['_oid'] = a.get(oid)
            registros.append(a)
        print(f'Descargados {len(registros)} de {len(ids)}')
    if len(registros) != len(ids) or len({r[oid] for r in registros}) != len(ids):
        raise RuntimeError('La descarga está incompleta o contiene IDs repetidos.')

    print('Campo de capacidad identificado:', C['capacidad'])
    if C['capacidad'] is None:
        print('ADVERTENCIA: no se encontró el campo de capacidad; revisa el listado de campos y alias.')
    filas = []
    for r in registros:
        lat, lon, alt = coordenadas(r, C)
        tipo = tipo_arroz(decodificar(valor(r, C['tipo']), C['tipo'], campos))
        c, ec = toneladas(valor(r, C['cant_c']), decodificar(valor(r, C['unidad_c']), C['unidad_c'], campos), valor(r, C['peso_c']), tipo in {'Cáscara', 'Ambas'})
        p, ep = toneladas(valor(r, C['cant_p']), decodificar(valor(r, C['unidad_p']), C['unidad_p'], campos), valor(r, C['peso_p']), tipo in {'Pilado', 'Ambas'})
        total = (c or 0) + (p or 0) if c is not None or p is not None else None
        filas.append({
            '_ObjectID': r[oid],
            'Provincia': valor(r, C['provincia']),
            'Cantón': valor(r, C['canton']),
            'Parroquia': valor(r, C['parroquia']),
            'Latitud': lat,
            'Longitud': lon,
            'Altura': alt,
            'Ruc': str(valor(r, C['ruc']) or '').strip(),
            'Razón Social': valor(r, C['razon']),
            'Nombre Apiladora': valor(r, C['piladora']),
            'Capacidad de arroz en Cáscara': valor(r, C['capacidad']),
            'Servicios': decodificar(valor(r, C['servicios']), C['servicios'], campos),
            'Tipo de Arroz': tipo,
            'Nombre Técnico': valor(r, C['tecnico']),
            'Cédula del Técnico': str(valor(r, C['cedula']) or '').strip(),
            'Fecha': fecha_ecuador(valor(r, C['fecha'])),
            'Observaciones': valor(r, C['observaciones']),
            'Arroz en cáscara (Tn)': c,
            'Arroz pilado (tn)': p,
            'Total almacenado (tn)': total,
            '_Estado': 'Correcto' if tipo != 'No identificado' and ec in {'Correcto', 'No aplica'} and ep in {'Correcto', 'No aplica'} else f'Cáscara: {ec}; Pilado: {ep}',
        })

    df = pd.DataFrame(filas)
    print('Valores no vacíos: latitud=', df['Latitud'].notna().sum(),
          'longitud=', df['Longitud'].notna().sum(),
          'altura=', df['Altura'].notna().sum(),
          'capacidad=', df['Capacidad de arroz en Cáscara'].notna().sum())
    # No deduplicar por RUC: un RUC puede representar varios establecimientos.
    # Orden cronológico de registros por RUC (solo informativo, no identifica establecimientos).
    df = df.sort_values(['Ruc', 'Fecha', '_ObjectID'], na_position='last').reset_index(drop=True)
    df['_Visita_RUC'] = df.groupby('Ruc').cumcount() + 1
    # Una sola hoja, exactamente las 19 columnas solicitadas.
    columnas = [
        'Provincia', 'Cantón', 'Parroquia', 'Latitud', 'Longitud', 'Altura',
        'Ruc', 'Razón Social', 'Nombre Apiladora',
        'Capacidad de arroz en Cáscara', 'Servicios', 'Tipo de Arroz',
        'Arroz en cáscara (Tn)', 'Arroz pilado (tn)', 'Total almacenado (tn)',
        'Nombre Técnico', 'Cédula del Técnico', 'Fecha', 'Observaciones'
    ]
    with pd.ExcelWriter(DESTINO, engine='openpyxl') as writer:
        df[columnas].to_excel(writer, sheet_name='REPORTE', index=False)
        ws = writer.sheets['REPORTE']
        ws.freeze_panes = 'D2'
        ws.auto_filter.ref = ws.dimensions
        for celda in ws[1]:
            celda.fill = PatternFill('solid', fgColor='17365D')
            celda.font = Font(bold=True, color='FFFFFF')
            celda.alignment = Alignment(horizontal='center', wrap_text=True)
        for j, nombre in enumerate(columnas, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(len(nombre) + 3, 17), 36)
            if nombre in ('Arroz en cáscara (Tn)', 'Arroz pilado (tn)', 'Total almacenado (tn)'):
                for i in range(2, ws.max_row + 1):
                    ws.cell(i, j).number_format = '#,##0.00'
            if nombre == 'Fecha':
                for i in range(2, ws.max_row + 1):
                    ws.cell(i, j).number_format = 'dd/mm/yyyy hh:mm'
            if nombre in ('Ruc', 'Cédula del Técnico'):
                for i in range(2, ws.max_row + 1):
                    ws.cell(i, j).number_format = '@'
    print('\nREPORTE:', DESTINO)
    print('Boletas:', len(df))
    print('RUC repetidos:', df.loc[df.duplicated('Ruc', keep=False) & df['Ruc'].ne(''), 'Ruc'].nunique())
    print('Posibles visitas posteriores por RUC (NO establecimientos confirmados):', int(((df['_Visita_RUC'] > 1) & df['Ruc'].ne('')).sum()))
    print('Registros con conversión pendiente:', int(df['_Estado'].ne('Correcto').sum()))
    print('Toneladas cáscara:', round(df['Arroz en cáscara (Tn)'].sum(), 2))
    print('Toneladas pilado:', round(df['Arroz pilado (tn)'].sum(), 2))
    if df['_Estado'].ne('Correcto').any():
        print('ATENCIÓN: existen registros pendientes de validar:')
        print(df.loc[df['_Estado'].ne('Correcto'), ['_ObjectID', '_Estado']].to_string(index=False))

if __name__ == '__main__':
    ejecutar()
